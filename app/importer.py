"""Importul Excel-ului cu tichete în SQLite.

Pași: citire foaie „Tichete” → completare severitate cu LLM (unde e goală)
→ clustere → salvare. Un import nou înlocuiește tot (tichete + rapoarte),
ceea ce e exact ce vrem în demo: re-import = pornim scenariul de la zero.
"""
import datetime as dt
import io
import logging

import openpyxl

from . import clustering, llm
from .config import DATA_DIR
from .db import tx
from .models import ImportResult

log = logging.getLogger("fixtrail.import")
LAST_IMPORT = DATA_DIR / "last_import.xlsx"

# antet din Excel → coloana din baza de date. Căutăm după nume, nu după poziție,
# ca o coloană mutată sau adăugată în Excel să nu strice importul.
COLUMNS = {
    "ID tichet": "id",
    "Data deschiderii": "data_deschiderii",
    "Canal": "canal",
    "Categorie": "categorie",
    "Funcționalitate": "functionalitate",
    "Titlu": "titlu",
    "Descriere": "descriere",
    "Pași de verificare": "pasi_verificare",
    "Severitate": "severitate",
    "Nr. raportări clienți": "nr_raportari",
    "Versiune raportată": "versiune_raportata",
    "Status": "status",
    "Versiune fix": "versiune_fix",
    "Data livrării fix-ului": "data_livrare_fix",
    "Echipă responsabilă": "echipa",
    "Tichet legat": "tichet_legat",
}
REQUIRED = {"id", "titlu", "status"}


def _clean(v):
    """Normalizează o celulă: date → ISO, text → fără spații, gol → None."""
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def parse_workbook(content: bytes) -> list[dict]:
    # data_only=True: citim valorile, nu formulele (ex. „Zile până la fix”)
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb["Tichete"] if "Tichete" in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h else "" for h in next(rows)]
    idx = {COLUMNS[h]: i for i, h in enumerate(header) if h in COLUMNS}
    missing = REQUIRED - idx.keys()
    if missing:
        raise ValueError(f"Lipsesc coloanele obligatorii: {', '.join(sorted(missing))}")

    tickets = []
    for n, row in enumerate(rows):
        t = {col: _clean(row[i]) if i < len(row) else None for col, i in idx.items()}
        if not t.get("id"):
            continue  # rânduri goale la final
        t["nr_raportari"] = int(t.get("nr_raportari") or 0)
        for k in ("versiune_raportata", "versiune_fix"):
            if t.get(k) is not None:
                t[k] = str(t[k])
        t["ordine"] = n
        tickets.append(t)
    return tickets


def import_excel(content: bytes) -> ImportResult:
    tickets = parse_workbook(content)
    if not tickets:
        raise ValueError("Foaia „Tichete” nu conține niciun tichet.")

    # 1. Severitate: unde e goală, decide LLM-ul (cu exemple din tichetele deja clasificate)
    empty = [t for t in tickets if not t.get("severitate")]
    examples = [t for t in tickets if t.get("severitate")]
    decided = llm.classify_severity(empty, examples)
    for t in tickets:
        t["severitate_sursa"] = "excel"
        if not t.get("severitate") and t["id"] in decided:
            t["severitate"] = decided[t["id"]]
            t["severitate_sursa"] = "llm"

    # 2. Clustere
    mapping, mode, acord = clustering.cluster(tickets)
    titles = {}
    for t in sorted(tickets, key=lambda x: x["id"]):  # titlul clusterului = al tichetului cel mai vechi
        titles.setdefault(mapping[t["id"]], t["titlu"])

    # 3. Salvare (înlocuim tot)
    cols = list(COLUMNS.values()) + ["severitate_sursa", "cluster_id", "ordine"]
    with tx() as c:
        c.execute("DELETE FROM tickets")
        c.execute("DELETE FROM reports")
        c.execute("DELETE FROM clusters")
        for t in tickets:
            t["cluster_id"] = mapping[t["id"]]
            c.execute(
                f"INSERT INTO tickets ({', '.join(cols)}) VALUES ({', '.join('?' * len(cols))})",
                [t.get(k) for k in cols],
            )
        c.executemany("INSERT INTO clusters(id, titlu) VALUES (?, ?)", list(titles.items()))

    # păstrăm fișierul ca să-l putem reîncărca dintr-un click (POST /api/import/reload)
    LAST_IMPORT.write_bytes(content)
    log.info("Import: %d tichete, %d clustere (%s)", len(tickets), len(titles), mode)
    return ImportResult(
        importate=len(tickets),
        severitatiCompletateDeAI=len(decided),
        clustere=len(titles),
        modClustere=mode,
        acordCuTichetLegat=acord,
    )


def reload_last() -> ImportResult:
    if not LAST_IMPORT.exists():
        raise FileNotFoundError("Nu există niciun import anterior.")
    return import_excel(LAST_IMPORT.read_bytes())


def import_file(path) -> ImportResult:
    """Import din linia de comandă / la pornire (vezi main.py)."""
    data = open(path, "rb").read()
    return import_excel(data)

